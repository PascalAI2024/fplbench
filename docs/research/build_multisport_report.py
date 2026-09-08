from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "research" / "multisport-benchmark-decision.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 28, 38)
MUTED = RGBColor(88, 96, 105)
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"
GREEN = RGBColor(39, 99, 71)
AMBER = RGBColor(122, 90, 0)
RED = RGBColor(155, 28, 28)


def set_run(run, size=11, bold=False, italic=False, color=INK, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        header = tr_pr.find(qn("w:tblHeader"))
        if row_index == 0 and header is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    set_run(p.add_run(str(text)), size=size, bold=bold, color=color)


def add_hyperlink(paragraph, text, url, size=9, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), str(color))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    r_pr.extend([fonts, color_node, underline, size_node])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run(run, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_numbering(doc, num_id, ordered=False):
    numbering = doc.part.numbering_part.element
    abs_id = str(50 + num_id)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), abs_id)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1." if ordered else "•")
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, fmt, text, jc, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_node = OxmlElement("w:abstractNumId")
    abs_node.set(qn("w:val"), abs_id)
    num.append(abs_node)
    numbering.append(num)


def add_list_item(doc, text, num_id=21):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, nid])
    p_pr.append(num_pr)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    set_run(p.add_run(text), size=11)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_callout(doc, label, text, color=GREEN):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F4F6F9")
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), str(color))
    borders.append(left)
    p_pr.extend([shading, borders])
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.08
    set_run(p.add_run(label + " "), size=11, bold=True, color=color)
    set_run(p.add_run(text), size=11)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    add_hyperlink(p, label, url, size=8.5)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 12, 6),
    ("Heading 2", 13, BLUE, 10, 5),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

add_numbering(doc, 21, ordered=False)
add_numbering(doc, 22, ordered=True)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
set_run(hp.add_run("FPLBENCH RESEARCH  |  MULTISPORT EXPANSION"), size=8.5, bold=True, color=MUTED)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_before = Pt(0)
set_run(fp.add_run("Decision memo  |  30 August 2026  |  Page "), size=8.5, color=MUTED)
add_page_field(fp)

# Memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(4)
set_run(p.add_run("MULTISPORT BENCHMARK EXPANSION"), size=23, bold=True, color=RGBColor(0, 0, 0))
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
set_run(p.add_run("Decision: which sport should receive the next leakage-safe public benchmark?"), size=14, color=RGBColor(55, 55, 55))

for label, value in (
    ("To", "FPLBench product and research owners"),
    ("From", "Multisport research coordinator"),
    ("Date", "30 August 2026"),
    ("Status", "Decision-ready; implementation not authorized by this memo"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(label + ": "), bold=True)
    set_run(p.add_run(value))

add_heading(doc, "Executive decision", 1)
add_callout(doc, "RECOMMENDATION", "Build NFLBench first. Build a separate soccer match benchmark second. Run cricket rights diligence before implementation.")
add_body(doc, "NFL is the closest extension of FPLBench: player-level forecasts, weekly cadence, canonical fantasy scoring, maintained CC-licensed historical data, and a clean pre-event freeze. The MVP should omit injury features unless a licensed timestamped source passes validation.")
add_body(doc, "Soccer is nearly as strong, but the clean product is match forecasting - home/draw/away probabilities and expected goals - rather than forcing another player-fantasy clone. Cricket is technically attractive, but match-archive reuse rights are not explicit enough for a public benchmark.")

add_heading(doc, "Ranked decision", 1)
rows = [
    ("1", "NFL player-week", "4.7", "GO", "Closest FPLBench extension; weekly and auditable"),
    ("2", "Soccer match", "4.6", "GO NEXT", "CC0 path; different target and product identity"),
    ("3", "T20 cricket", "3.6", "CONDITIONAL", "Excellent corpus; archive rights and fixtures unresolved"),
    ("4", "MLB player/game", "3.4", "DELAYED", "Permissive history; live rights and daily churn"),
    ("5", "NHL daily fantasy", "2.8", "HOLD", "Technical API access without reuse license"),
    ("6", "NBA daily fantasy", "2.6", "NO-GO", "Official terms restrict fantasy/database use"),
    ("7", "Formula 1 fantasy", "2.5", "HOLD", "Rights and official-label completeness"),
]
table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
headers = ["Rank", "Candidate", "Score", "Verdict", "Decision reason"]
for i, value in enumerate(headers):
    set_cell_text(table.rows[0].cells[i], value, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(table.rows[0].cells[i], PALE_BLUE)
for row in rows:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
        color = GREEN if value in ("GO", "GO NEXT") else AMBER if value in ("CONDITIONAL", "DELAYED", "HOLD") else RED if value == "NO-GO" else INK
        set_cell_text(cells[i], value, bold=i in (0, 3), size=8.7, align=align, color=color)
        if len(table.rows) % 2 == 1:
            set_cell_shading(cells[i], "FAFBFC")
set_table_geometry(table, [600, 1940, 750, 1250, 4820])
p = doc.add_paragraph("Decision scores use these weights: data rights/source stability 25%; leakage control 20%; cadence 15%; scoring clarity 15%; audience/product fit 15%; engineering effort 10%.")
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(4)
set_run(p.runs[0], size=8.5, italic=True, color=MUTED)

add_heading(doc, "Why NFL is first", 1)
add_heading(doc, "Operating fit and sources", 2)
add_body(doc, "NFL has a natural weekly forecasting round. That leaves enough time to freeze, publish, inspect, and score immutable artifacts without the daily lineup churn of MLB, NBA, or NHL.")
add_body(doc, "nflverse publishes automated releases under CC BY 4.0. Its update schedule says player statistics refresh nightly, rosters daily, schedules every five minutes, and depth charts carry ISO timestamps from 2025. Thursday is the cleanest post-game scoring point after corrections. The same source states that its injury feed ended after 2024.")
add_body(doc, "Sleeper's read-only API exposes NFL league settings, rosters, matchups, player identities, practice participation, injury status, and update metadata. It is free for non-commercial use; commercial use requires licensing. The player map is intended to be cached and called no more than once daily.")
add_heading(doc, "Rights boundary", 2)
add_callout(doc, "DO NOT SCRAPE NFL.COM", "NFL.com restricts systematic retrieval and database construction without consent. Official injury pages define the public reporting boundary, but they are not an authorized ingestion feed.", color=RED)
add_body(doc, "Version one should therefore use no injury features unless Sleeper's licensed metadata is sufficient for the research use and its timestamp semantics pass validation. A missing feature is preferable to an untracked rights problem.")

add_heading(doc, "NFLBench MVP contract", 1)
add_heading(doc, "Forecast unit, freeze, and target", 2)
for item in (
    "One row per season, week, player, team, opponent, and frozen snapshot; positions QB, RB, WR, and TE.",
    "Player universe comes from the prior roster plus the last eligible timestamped depth-chart snapshot. Frozen DNPs remain and score zero.",
    "Primary track freezes at least 24 hours before the first weekly kickoff. A later T-120-minute track, if added, gets a separate leaderboard.",
    "Fixed PPR target: passing TD 4; rushing/receiving TD 6; passing yards 1/25; rushing/receiving yards 1/10; reception 1; interception -2; lost fumble -2.",
    "Score after the Thursday post-correction refresh and record the exact nflverse release or content hash.",
):
    add_list_item(doc, item, 21)

add_heading(doc, "Evaluation", 2)
add_body(doc, "Primary metric is mean absolute error over all frozen player-weeks, including zeros and DNPs. Secondary diagnostics are RMSE, median absolute error, interval coverage when probabilistic outputs exist, and Spearman or NDCG@20 for rank quality.")
add_body(doc, "Baselines are prior-four-week exponentially weighted mean and position median. External consensus projections belong only when historical snapshot and redistribution rights are explicit. Suggested validation is 2016-2024 training, 2025 locked validation, then a six-to-eight-week 2026 live pilot.")

add_heading(doc, "Leakage controls", 2)
for item in (
    "Every feature row carries observed_at, source_revision, and ingested_at.",
    "Feature joins enforce observed_at <= cutoff_at; same-game statistics are strictly unavailable to that game's forecast.",
    "Forecast CSV, source manifest, configuration, and model revision are hashed and published before kickoff.",
    "Late inactive news, corrections, depth-chart changes, and injuries never rewrite a frozen artifact.",
    "Targets are computed in a separate scoring job after the correction window.",
    "Identity joins fail closed on missing or ambiguous player mappings.",
):
    add_list_item(doc, item, 22)

add_heading(doc, "Acceptance gates", 2)
for item in (
    "A legal/source manifest names each dataset, license, attribution, cache rule, and commercial limitation.",
    "A synthetic time-travel test proves post-cutoff records cannot enter a feature row.",
    "A DNP fixture proves frozen players remain and score zero.",
    "A stat-correction fixture proves forecasts stay immutable while targets can be versioned.",
    "A clean-checkout dry run reproduces one published week and matches hashes.",
    "The public board shows cutoff, artifact hash, source revisions, forecasts, realized points, and baselines.",
):
    add_list_item(doc, item, 21)

add_heading(doc, "Second product: SoccerBench", 1)
add_body(doc, "OpenFootball publishes current multi-league fixtures and results and dedicates its data and schema to the public domain. football-data.org can supplement a limited competition set under attribution, rate, and application restrictions.")
add_body(doc, "Freeze at T-24h. Predict home/draw/away probabilities and expected home and away goals. Use multiclass log loss as primary scoring, with Brier score, calibration, and goal MAE or Poisson deviance as secondary metrics. Avoid UEFA Fantasy as an operational feed because UEFA terms restrict systematic collection, database construction, scraping, and model development.")
add_callout(doc, "PRODUCT BOUNDARY", "SoccerBench needs its own identity and leaderboard. Combining match-probability scores with player-fantasy MAE would be numerically tidy and conceptually useless.", color=DARK_BLUE)

add_heading(doc, "Conditional third option: T20 CricketBench", 1)
add_body(doc, "Cricsheet offers more than 22,000 ball-by-ball matches, while its Register maps player identifiers under ODC Attribution 1.0. The technical corpus is excellent. The public-license gap is material: the explicit Register license does not clearly grant reuse of the match archives, and a licensed live fixture/roster source has not been verified.")
add_body(doc, "If cleared, start match-level at T-60m before the toss: winner probability and expected innings totals. Score with log loss/Brier and MAE or CRPS. Exclude playing XI, toss, DLS revisions, impact-player state, and post-start corrections from the frozen track.")

add_heading(doc, "Why the other sports wait", 1)
for title, text in (
    ("MLB", "Retrosheet permits redistribution and commercial reuse with attribution and supplies complete AL/NL play-by-play through 2025. MLB.com restricts automated collection, while daily pitchers and lineups move close to game time. Use a historical or annually delayed model unless a licensed current feed is obtained."),
    ("NHL", "Official JSON endpoints expose useful data and NHL Fantasy Stars defines a daily target, but NHL terms restrict scraping, database entry, distribution, and derivative reuse. An endpoint is not a license."),
    ("NBA", "NBA terms expressly restrict using NBA Statistics with a fantasy game or comprehensive regularly updated database without consent. Game-day injury updates and conflicting official fantasy formulas add operational ambiguity."),
    ("Formula 1", "F1 claims exclusive rights over results, timing, and statistics and prohibits text/data mining or AI use without permission. Jolpica is non-commercial; OpenF1 is unofficial/personal-use and incomplete for exact official Fantasy scoring."),
):
    add_heading(doc, title, 2)
    add_body(doc, text)

add_heading(doc, "Recommended sequence", 1)
for item in (
    "Two-week NFLBench feasibility spike: source/license manifest, point-in-time schema, one historical reconstruction, and one immutable sample board.",
    "Six-to-eight-week live NFL pilot: no injury features, T-24h track only, public baselines, correction-aware scoring.",
    "SoccerBench prototype: OpenFootball-only, fixed competitions, match probabilities, and expected goals.",
    "Cricket diligence: written archive-license confirmation plus live fixture/roster provenance; build only after both pass.",
):
    add_list_item(doc, item, 22)

add_heading(doc, "Confidence and unresolved items", 1)
add_body(doc, "High confidence: NFL ranks first; an OpenFootball-backed soccer match benchmark is viable; NBA/NHL/F1 rights concerns are material; Retrosheet supports a historical MLB product.")
add_body(doc, "Medium confidence: the commercial path for Sleeper metadata, depth-chart timestamp consistency before 2025, and cricket match-archive reuse rights. Before expansion beyond research, obtain a narrow rights review and written confirmation for sources whose terms do not explicitly cover redistribution.")

add_heading(doc, "Primary source register", 1)
sources = [
    ("nflverse data releases", "https://github.com/nflverse/nflverse-data"),
    ("nflverse update schedule", "https://nflreadr.nflverse.com/articles/nflverse_data_schedule.html"),
    ("nflreadpy documentation", "https://github.com/nflverse/nflreadpy/blob/main/docs/index.md"),
    ("Sleeper API", "https://docs.sleeper.com/"),
    ("NFL injury reporting cadence", "https://operations.nfl.com/calendar-events/nfl-important-dates"),
    ("NFL.com terms", "https://www.nfl.com/legal/terms/"),
    ("NFL Fantasy official rules", "https://static.www.nfl.com/image/upload/v1745955215/league/apps/fantasy/media/rules/OfficialRules2025.pdf"),
    ("OpenFootball football.json", "https://github.com/openfootball/football.json"),
    ("football-data.org pricing", "https://www.football-data.org/pricing"),
    ("Hudl/StatsBomb open data", "https://github.com/hudl/open-data"),
    ("UEFA terms", "https://www.uefa.com/termsconditions/"),
    ("Cricsheet downloads", "https://cricsheet.org/downloads/"),
    ("Cricsheet Register", "https://cricsheet.org/register/"),
    ("ICC website terms", "https://www.icc-cricket.com/about/the-icc/legal-notices/website-terms-of-use"),
    ("Retrosheet data-use policy", "https://www.retrosheet.org/datause.html"),
    ("MLB terms", "https://www.mlb.com/official-information/terms-of-use"),
    ("NHL terms", "https://www.nhl.com/info/terms-of-service"),
    ("NBA terms", "https://www.nba.com/termsofuse"),
    ("Formula 1 guidelines", "https://www.formula1.com/en/information/guidelines.4EOKE9RRqevL4niTK9kWyt"),
    ("Jolpica terms", "https://github.com/jolpica/jolpica-f1/blob/main/TERMS.md"),
    ("OpenF1 documentation", "https://openf1.org/docs/"),
]
for label, url in sources:
    add_source(doc, label, url)

doc.core_properties.title = "Multisport Benchmark Expansion"
doc.core_properties.subject = "Decision memo for the next leakage-safe FPLBench-style sports benchmark"
doc.core_properties.author = "FPLBench Research"
doc.core_properties.keywords = "sports forecasting, benchmark, leakage, NFL, soccer, cricket"
doc.core_properties.comments = "Generated from multisport-benchmark-decision-source.md"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
